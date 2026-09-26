from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from html.parser import HTMLParser
from io import BytesIO
import csv
import ipaddress
import json
import logging
import re
import socket
import time
from pathlib import Path
from urllib.parse import urljoin, urlsplit
from uuid import uuid4

import httpx
import pandas as pd
from PIL import Image, ImageDraw, ImageFont, ImageOps, UnidentifiedImageError

from agents.price_source_agent import run_price_source_agent
from agents.schemas.price_source_schema import (
    CANONICAL_UNIT_CODES,
    PRICE_SOURCE_CATEGORIES,
)
from db.company_access import assert_company_owner
from db.repositories import insert_agent_usage_event
from db.supabase_client import get_supabase_client


PRICE_SOURCE_BUCKET = "company-price-sources"
MAX_SOURCE_BYTES = 50 * 1024 * 1024
LEGACY_EXTREME_RATIO = 3.0
LEGACY_CONFIDENCE_PENALTY = 15.0
SUPPORTED_SUFFIXES = {".pdf", ".xlsx", ".csv", ".jpg", ".jpeg", ".png"}
CONTENT_TYPES = {
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
}

PRICE_SOURCE_DEPARTMENTS = ("Wood", "Metal", "Finishing")
PRICE_CATALOG_DEPARTMENTS = {
    "Wood Sheets": "Wood",
    "Solid Wood": "Wood",
    "Wood Supplies": "Wood",
    "Glass": "Wood",
    "Metal Sheets": "Metal",
    "Metal Profiles": "Metal",
    "Metal Supplies": "Metal",
    "Metal": "Metal",
    "Paints & Coatings": "Finishing",
    "Coating Supplies": "Finishing",
    "Other": "Wood",
}
LEGACY_PRICE_SOURCE_CATEGORIES = {
    "Sheet Materials": "Wood Sheets",
    "Hardware": "Wood Supplies",
    "Edgebanding": "Wood Supplies",
    "Adhesives and Consumables": "Wood Supplies",
    "Finishes and Coatings": "Paints & Coatings",
    "Abrasives and Sanding": "Coating Supplies",
}
PRICE_CATALOG_DEPARTMENT_ORDER = {"Wood": 0, "Metal": 1, "Finishing": 2}


class PriceSourceError(ValueError):
    pass


@dataclass(frozen=True)
class CombinedPriceSource:
    name: str
    data: bytes

    def getvalue(self) -> bytes:
        return self.data


@dataclass(frozen=True)
class PriceSourceProcessResult:
    source_id: str
    summary: dict[str, object]


logger = logging.getLogger(__name__)


def validate_price_source_upload_selection(files: list) -> None:
    """Reject multi-file selections that cannot represent one logical document."""
    selected = [item for item in files if item is not None]
    if len(selected) <= 1:
        return
    suffixes = [Path(str(item.name)).suffix.lower() for item in selected]
    if any(suffix not in {".jpg", ".jpeg", ".png"} for suffix in suffixes):
        raise PriceSourceError(
            "Select one PDF or spreadsheet, or several JPEG/PNG photos from the same document"
        )


def accepted_price_source_uploads(files: list) -> list:
    """Apply the one-document MVP contract to a native multi-file selection."""
    selected = [item for item in files if item is not None]
    if len(selected) <= 1:
        return selected
    suffixes = [Path(str(item.name)).suffix.lower() for item in selected]
    photo_suffixes = {".jpg", ".jpeg", ".png"}
    if suffixes[0] in photo_suffixes:
        return [
            item
            for item, suffix in zip(selected, suffixes, strict=True)
            if suffix in photo_suffixes
        ]
    return selected[:1]


def render_price_source_pdf_preview(
    file_bytes: bytes,
    *,
    max_width: int = 240,
    max_height: int = 140,
) -> bytes | None:
    """Render a small first-page PNG without affecting source acceptance."""
    if not file_bytes:
        return None
    try:
        import pymupdf

        document = pymupdf.open(stream=file_bytes, filetype="pdf")
        try:
            if document.needs_pass or document.page_count < 1:
                return None
            page = document[0]
            width = max(float(page.rect.width), 1.0)
            height = max(float(page.rect.height), 1.0)
            scale = max(0.1, min(2.0, max_width / width, max_height / height))
            pixmap = page.get_pixmap(
                matrix=pymupdf.Matrix(scale, scale),
                colorspace=pymupdf.csRGB,
                alpha=False,
            )
            return pixmap.tobytes("png")
        finally:
            document.close()
    except Exception:
        logger.info("Price source PDF preview unavailable", exc_info=True)
        return None


def _price_source_preview_font(size: int) -> ImageFont.ImageFont:
    for path in (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ):
        try:
            return ImageFont.truetype(path, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _render_price_source_table_preview(rows: list[list[object]]) -> bytes | None:
    if not rows:
        return None
    width, height = 240, 140
    visible_rows = rows[:7]
    column_count = min(max((len(row) for row in visible_rows), default=0), 6)
    if column_count < 1:
        return None
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font = _price_source_preview_font(10)
    row_height = height / len(visible_rows)
    column_width = width / column_count
    for row_index, row in enumerate(visible_rows):
        top = round(row_index * row_height)
        bottom = round((row_index + 1) * row_height)
        fill = "#F1EBFA" if row_index == 0 else ("#FFFFFF" if row_index % 2 else "#FAF8FC")
        draw.rectangle((0, top, width, bottom), fill=fill)
        for column_index in range(column_count):
            left = round(column_index * column_width)
            right = round((column_index + 1) * column_width)
            draw.rectangle((left, top, right, bottom), outline="#DED8E5", width=1)
            value = (
                ""
                if column_index >= len(row) or row[column_index] is None
                else str(row[column_index])
            )
            value = value.replace("\n", " ").strip()
            max_chars = max(3, int(column_width / 6.2) - 1)
            if len(value) > max_chars:
                value = value[: max_chars - 1] + "…"
            draw.text((left + 3, top + 3), value, font=font, fill="#302A36")
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    image.close()
    return output.getvalue()


def render_price_source_preview(
    file_name: str,
    file_bytes: bytes,
    *,
    max_width: int = 240,
    max_height: int = 140,
) -> bytes | None:
    """Render a bounded preview for every accepted Price Source format."""
    suffix = Path(str(file_name)).suffix.lower()
    if not file_bytes:
        return None
    if suffix == ".pdf":
        return render_price_source_pdf_preview(
            file_bytes,
            max_width=max_width,
            max_height=max_height,
        )
    if suffix in {".jpg", ".jpeg", ".png"}:
        try:
            with Image.open(BytesIO(file_bytes)) as source:
                image = ImageOps.exif_transpose(source).convert("RGB")
                image.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)
                output = BytesIO()
                image.save(output, format="PNG", optimize=True)
                image.close()
                return output.getvalue()
        except (UnidentifiedImageError, OSError):
            logger.info("Price source image preview unavailable", exc_info=True)
            return None
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
            try:
                worksheet = next(
                    (sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"),
                    workbook.worksheets[0],
                )
                rows = [list(row[:6]) for row in worksheet.iter_rows(max_row=7, values_only=True)]
            finally:
                workbook.close()
            return _render_price_source_table_preview(rows)
        except Exception:
            logger.info("Price source spreadsheet preview unavailable", exc_info=True)
            return None
    if suffix == ".csv":
        try:
            text = file_bytes.decode("utf-8-sig", errors="replace")
            rows = [row[:6] for _, row in zip(range(7), csv.reader(text.splitlines()))]
            return _render_price_source_table_preview(rows)
        except Exception:
            logger.info("Price source CSV preview unavailable", exc_info=True)
            return None
    return None


def combine_price_source_files(files: list) -> object | None:
    """Keep one upload as-is or combine ordered JPEG/PNG pages into one PDF."""
    selected = accepted_price_source_uploads(files)
    if not selected:
        return None
    if len(selected) == 1:
        return selected[0]
    validate_price_source_upload_selection(selected)
    if sum(len(item.getvalue()) for item in selected) > MAX_SOURCE_BYTES:
        raise PriceSourceError("The combined price source must be 50 MB or smaller")
    pages: list[Image.Image] = []
    try:
        for item in selected:
            with Image.open(BytesIO(item.getvalue())) as image:
                page = image.convert("RGB")
                page.load()
                pages.append(page)
        output = BytesIO()
        pages[0].save(output, format="PDF", save_all=True, append_images=pages[1:])
    except (UnidentifiedImageError, OSError) as exc:
        raise PriceSourceError("One of the selected photos could not be read") from exc
    finally:
        for page in pages:
            page.close()
    return CombinedPriceSource(
        name=f"photo-document-{len(selected)}-pages.pdf",
        data=output.getvalue(),
    )


def _emit_duration(trace, name: str, started_at: float, **metadata: object) -> None:
    if trace is not None:
        trace.event(
            name,
            duration_ms=(time.perf_counter() - started_at) * 1000,
            metadata=metadata,
        )


class _VisibleTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._ignored = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"}:
            self._ignored += 1
        if tag.lower() in {"p", "div", "tr", "li", "br", "h1", "h2", "h3", "td", "th"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript", "svg"} and self._ignored:
            self._ignored -= 1
        if tag.lower() in {"p", "div", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._ignored:
            self.parts.append(data)

    def text(self) -> str:
        return re.sub(r"[ \t]+", " ", "".join(self.parts)).strip()


def _normalized_name(value: str) -> str:
    return re.sub(r"[^\w]+", " ", value.casefold(), flags=re.UNICODE).strip()


def _normalized_unit(value: str) -> str:
    unit = _normalized_name(value).replace(" ", "")
    aliases = {
        "sqm": "m2",
        "m²": "m2",
        "squaremeter": "m2",
        "squaremeters": "m2",
        "m2": "m2",
        "sqmeter": "m2",
        "sqmeters": "m2",
        "squaremetre": "m2",
        "squaremetres": "m2",
        "lm": "linear_m",
        "linm": "linear_m",
        "runningmeter": "linear_m",
        "runningmeters": "linear_m",
        "linearmeter": "linear_m",
        "linearmeters": "linear_m",
        "unit": "piece",
        "units": "piece",
        "pcs": "piece",
        "pc": "piece",
        "sheet": "sheet",
        "sheets": "sheet",
    }
    return aliases.get(unit, unit)


def _decimal_equal(left: object, right: object, places: str) -> bool:
    try:
        quantum = Decimal(places)
        return Decimal(str(left)).quantize(quantum) == Decimal(str(right)).quantize(quantum)
    except (InvalidOperation, TypeError, ValueError):
        return False


def price_offer_matches_row(
    offer: dict,
    row: dict,
    *,
    default_currency: str = "",
) -> bool:
    """Compare persisted price-affecting values at their database precision."""
    currency = str(row.get("raw_currency") or default_currency or "").strip().upper()
    vat_included = (
        True
        if row.get("raw_vat_mode") == "included"
        else False if row.get("raw_vat_mode") == "excluded" else None
    )
    return all(
        (
            _decimal_equal(offer.get("source_price"), row.get("raw_price"), "0.0001"),
            _normalized_unit(str(offer.get("source_unit") or ""))
            == _normalized_unit(str(row.get("raw_unit") or "")),
            _normalized_unit(str(offer.get("purchase_unit") or ""))
            == _normalized_unit(str(row.get("purchase_unit") or "")),
            _normalized_unit(str(offer.get("calculation_unit") or ""))
            == _normalized_unit(str(row.get("calculation_unit") or "")),
            _decimal_equal(
                offer.get("conversion_factor"), row.get("conversion_factor"), "0.00000001"
            ),
            _decimal_equal(
                offer.get("normalized_price"), row.get("normalized_price"), "0.0001"
            ),
            str(offer.get("currency") or "").strip().upper() == currency,
            offer.get("vat_included") is vat_included,
        )
    )


def apply_legacy_price_benchmark(result: dict, legacy_materials: list[dict]) -> dict:
    """Use exact legacy identity/unit matches only as a negative confidence signal."""
    benchmark: dict[tuple[str, str], float] = {}
    for material in legacy_materials:
        name = _normalized_name(str(material.get("material_name") or ""))
        unit = _normalized_unit(str(material.get("price_unit") or ""))
        price = material.get("price")
        if name and unit and isinstance(price, (int, float)) and price > 0:
            benchmark[(name, unit)] = float(price)

    for row in result.get("rows") or []:
        if row.get("status") != "ready":
            continue
        key = (
            _normalized_name(str(row.get("normalized_name") or "")),
            _normalized_unit(str(row.get("calculation_unit") or "")),
        )
        old_price = benchmark.get(key)
        new_price = float(row.get("normalized_price") or 0)
        if not old_price or new_price <= 0:
            continue
        ratio = new_price / old_price
        if ratio > LEGACY_EXTREME_RATIO or ratio < 1 / LEGACY_EXTREME_RATIO:
            row["confidence"] = max(0.0, float(row["confidence"]) - LEGACY_CONFIDENCE_PENALTY)
            row["reason_codes"] = sorted(
                set((row.get("reason_codes") or []) + ["extreme_legacy_price_difference"])
            )
    return result


def canonical_price_source_category(category: str) -> str:
    normalized = category.strip() or "Other"
    return LEGACY_PRICE_SOURCE_CATEGORIES.get(normalized, normalized)


def price_source_material_types(result: dict) -> list[str]:
    return sorted(
        {
            canonical_price_source_category(str(row.get("material_type") or "Other"))
            for row in result.get("rows") or []
            if row.get("status") != "excluded"
        },
        key=str.casefold,
    )


def guard_price_source_department(result: dict, department: str) -> dict:
    """Keep rows outside an explicitly selected department non-active."""
    if not department:
        return result
    for row in result.get("rows") or []:
        if not isinstance(row, dict) or row.get("status") == "excluded":
            continue
        material_type = canonical_price_source_category(
            str(row.get("material_type") or "Other")
        )
        detected_department = PRICE_CATALOG_DEPARTMENTS.get(material_type)
        if material_type != "Other" and detected_department != department:
            row["status"] = "unresolved"
            row["reason_codes"] = sorted(
                set((row.get("reason_codes") or []) + ["selected_department_mismatch"])
            )
    return result


def price_source_semantic_fingerprint(result: dict) -> str:
    """Identify the same commercial document across files, scans, or photos."""
    document_number = _normalized_name(str(result.get("document_number") or ""))
    basis: dict[str, object] = {
        "source_origin": str(result.get("source_origin") or "unknown"),
        "supplier": _normalized_name(str(result.get("supplier_name") or "")),
        "document_number": document_number,
        "document_date": str(result.get("document_date") or ""),
    }
    if not document_number:
        basis.update(
            {
                "document_type": str(result.get("document_type") or ""),
                "currency": str(result.get("currency") or ""),
                "document_total": float(result.get("document_total") or 0),
            }
        )
        basis["rows"] = [
            {
                "description": _normalized_name(str(row.get("raw_description") or "")),
                "sku": _normalized_name(str(row.get("raw_sku") or "")),
                "price": float(row.get("raw_price") or 0),
                "quantity": float(row.get("raw_quantity") or 0),
                "line_total": float(row.get("raw_line_total") or 0),
            }
            for row in result.get("rows") or []
            if row.get("status") != "excluded"
        ]
    encoded = json.dumps(basis, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return sha256(encoded.encode("utf-8")).hexdigest()


def price_source_template_fingerprint(file_name: str, data: bytes) -> str | None:
    """Identify a recurring spreadsheet layout while ignoring row values."""
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        rows = list(csv.reader(data.decode("utf-8-sig", errors="replace").splitlines()))
        if not rows:
            return None
        basis = {
            "format": "csv",
            "columns": len(rows[0]),
            "headers": [_normalized_name(cell) for cell in rows[0]],
        }
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(data), read_only=False, data_only=False)
        except Exception:
            return None
        try:
            sheets: list[dict[str, object]] = []
            cell_reference = re.compile(r"(?<![A-Z0-9_])(\$?[A-Z]{1,3})\$?\d+")
            for sheet in workbook.worksheets:
                row_shapes: list[tuple] = []
                for row in sheet.iter_rows(
                    min_row=1,
                    max_row=min(sheet.max_row, 250),
                    min_col=1,
                    max_col=min(sheet.max_column, 80),
                ):
                    cells: list[tuple] = []
                    for cell in row:
                        value = cell.value
                        if value is None:
                            continue
                        if cell.data_type == "f":
                            kind = "formula"
                            anchor = cell_reference.sub(r"\1#", str(value))
                        elif isinstance(value, bool):
                            kind = "bool"
                            anchor = ""
                        elif isinstance(value, (int, float, Decimal)):
                            kind = "number"
                            anchor = ""
                        elif getattr(cell, "is_date", False):
                            kind = "date"
                            anchor = ""
                        else:
                            kind = "text"
                            anchor = (
                                _normalized_name(str(value))[:80]
                                if cell.row <= 3 or bool(cell.font and cell.font.bold)
                                else ""
                            )
                        style = (
                            bool(cell.font and cell.font.bold),
                            str(cell.number_format or ""),
                            str(cell.alignment.horizontal or ""),
                            str(cell.fill.fill_type or ""),
                        )
                        cells.append((cell.column, kind, anchor, style))
                    shape = tuple(cells)
                    if shape and (not row_shapes or row_shapes[-1] != shape):
                        row_shapes.append(shape)
                sheets.append(
                    {
                        "title": _normalized_name(sheet.title),
                        "state": sheet.sheet_state,
                        "merged": sorted(str(item) for item in sheet.merged_cells.ranges),
                        "row_shapes": row_shapes,
                    }
                )
            basis = {"format": "xlsx", "sheets": sheets}
        finally:
            workbook.close()
    else:
        return None
    encoded = json.dumps(basis, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return sha256(encoded.encode("utf-8")).hexdigest()


def price_source_family_identity(
    result: dict,
    *,
    source_name: str,
    source_kind: str,
    template_sha256: str | None = None,
) -> tuple[str, str]:
    """Return a stable internal identity for successive revisions of one source."""
    if source_kind == "url":
        parsed = urlsplit(source_name)
        origin = f"{(parsed.hostname or '').casefold()}{parsed.path.rstrip('/').casefold()}"
    elif template_sha256:
        origin = f"template:{template_sha256}"
    else:
        origin = _normalized_name(Path(source_name).name)
    basis: dict[str, str] = {
        "source_origin": _normalized_name(str(result.get("source_origin") or "unknown")),
        "supplier": _normalized_name(str(result.get("supplier_name") or "")),
        "source_kind": source_kind,
        "origin": origin,
    }
    if not template_sha256:
        basis["document_type"] = _normalized_name(
            str(result.get("document_type") or "")
        )
    if origin.startswith("photo document"):
        basis["document_number"] = _normalized_name(
            str(result.get("document_number") or "")
        )
        basis["document_date"] = str(result.get("document_date") or "")
    encoded = json.dumps(basis, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    fingerprint = sha256(encoded.encode("utf-8")).hexdigest()
    return fingerprint, f"PSF-{fingerprint[:12].upper()}"


def _find_previous_source_revision(
    client,
    company_id: str,
    *,
    family_sha256: str,
    semantic_sha256: str,
) -> tuple[dict | None, int]:
    sources = (
        client.table("company_price_sources")
        .select("source_id,processing_summary,created_at")
        .eq("company_id", company_id)
        .neq("status", "archived")
        .order("created_at", desc=True)
        .execute()
    ).data or []
    family_matches = [
        source
        for source in sources
        if isinstance(source.get("processing_summary"), dict)
        and source["processing_summary"].get("source_family_sha256") == family_sha256
    ]
    if family_matches:
        revision = max(
            int((source.get("processing_summary") or {}).get("source_revision") or 1)
            for source in family_matches
        ) + 1
        return family_matches[0], revision
    semantic_match = next(
        (
            source
            for source in sources
            if isinstance(source.get("processing_summary"), dict)
            and source["processing_summary"].get("semantic_sha256") == semantic_sha256
        ),
        None,
    )
    return semantic_match, 2 if semantic_match else 1


def _unchanged_duplicate_summary(source: dict) -> dict[str, object]:
    previous = dict(source.get("processing_summary") or {})
    total = int(previous.get("total") or 0)
    ready = int(previous.get("ready") or 0)
    summary = {
        "total": total,
        "ready": ready,
        "new": 0,
        "updated": 0,
        "unchanged": ready,
        "unresolved": int(previous.get("unresolved") or 0),
        "excluded": int(previous.get("excluded") or 0),
        "exact_duplicate": True,
        "agent_duration_seconds": 0.0,
        "token_cost": 0.0,
    }
    for key in (
        "source_family_sha256",
        "source_family_code",
        "template_sha256",
        "source_revision",
        "previous_source_id",
    ):
        if previous.get(key) is not None:
            summary[key] = previous[key]
    return summary


def _validate_department(department: str) -> str:
    normalized = department.strip()
    if normalized and normalized not in PRICE_SOURCE_DEPARTMENTS:
        raise PriceSourceError("Choose a department.")
    return normalized


def _validate_public_url(url: str) -> str:
    candidate = url.strip()
    parsed = urlsplit(candidate)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise PriceSourceError("Enter a public HTTP or HTTPS URL.")
    if parsed.username or parsed.password or parsed.port not in {None, 80, 443}:
        raise PriceSourceError("The URL must not contain credentials or a custom port.")
    try:
        addresses = {
            item[4][0]
            for item in socket.getaddrinfo(parsed.hostname, parsed.port or 443, type=socket.SOCK_STREAM)
        }
    except OSError as exc:
        raise PriceSourceError("The supplier page could not be resolved.") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if not ip.is_global:
            raise PriceSourceError("Only public supplier pages can be imported.")
    return candidate


def fetch_public_page(url: str, *, client: httpx.Client | None = None) -> tuple[str, bytes, str]:
    current = _validate_public_url(url)
    http = client or httpx.Client(timeout=20, headers={"User-Agent": "CosterlyAIPriceImporter/1.0"})
    owns_client = client is None
    try:
        for _ in range(4):
            response = http.get(current, follow_redirects=False)
            if response.status_code in {301, 302, 303, 307, 308}:
                location = response.headers.get("location")
                if not location:
                    raise PriceSourceError("The supplier page returned an invalid redirect.")
                current = _validate_public_url(urljoin(current, location))
                continue
            response.raise_for_status()
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            if content_type not in {"text/html", "text/plain"}:
                raise PriceSourceError("The URL must point to a readable supplier webpage.")
            content = response.content
            if len(content) > MAX_SOURCE_BYTES:
                raise PriceSourceError("The supplier page is too large to process.")
            parser = _VisibleTextParser()
            parser.feed(response.text)
            visible_text = parser.text()
            if not visible_text:
                raise PriceSourceError(
                    "This supplier page does not expose readable text. "
                    "Upload its PDF, screenshot, or photo instead."
                )
            return current, content, visible_text
        raise PriceSourceError("The supplier page redirected too many times.")
    except httpx.HTTPError as exc:
        raise PriceSourceError("The supplier page could not be downloaded.") from exc
    finally:
        if owns_client:
            http.close()


def extract_spreadsheet_text(file_name: str, data: bytes) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(text.splitlines()))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)[:180_000]
    if suffix == ".xlsx":
        sheets = pd.read_excel(BytesIO(data), sheet_name=None, dtype=str)
        parts: list[str] = []
        for sheet_name, frame in sheets.items():
            parts.append(f"SHEET: {sheet_name}")
            parts.append(frame.fillna("").to_csv(index=False))
        return "\n".join(parts)[:180_000]
    return ""


def _source_extension(file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise PriceSourceError("Upload PDF, XLSX, CSV, JPEG, or PNG")
    return suffix


def _date_or_none(value: str):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date().isoformat() if value else None
    except ValueError:
        return None


def list_price_sources(access) -> list[dict]:
    client = get_supabase_client()
    assert_company_owner(client, str(access.user_id), str(access.company_id))
    return (
        client.table("company_price_sources")
        .select(
            "source_id,source_name,source_kind,source_url,storage_path,mime_type,"
            "category,document_type,document_date,currency,vat_mode,status,"
            "processing_summary,created_at,processed_at,company_suppliers(supplier_name)"
        )
        .eq("company_id", str(access.company_id))
        .neq("status", "archived")
        .order("created_at", desc=True)
        .execute()
    ).data or []


def list_price_catalog(access) -> list[dict]:
    """Return active normalized supplier offers enriched for catalog display."""
    client = get_supabase_client()
    company_id = str(access.company_id)
    assert_company_owner(client, str(access.user_id), company_id)
    offers = (
        client.table("company_material_offers")
        .select(
            "offer_id,company_material_id,supplier_id,source_id,source_row_id,"
            "source_price,source_unit,purchase_unit,calculation_unit,conversion_factor,"
            "normalized_price,normalized_unit,currency,vat_included,valid_from,"
            "confidence,created_at"
        )
        .eq("company_id", company_id)
        .eq("status", "active")
        .execute()
    ).data or []
    if not offers:
        return []

    materials = (
        client.table("company_material_items")
        .select("company_material_id,category,canonical_name,preferred_unit")
        .eq("company_id", company_id)
        .neq("status", "archived")
        .execute()
    ).data or []
    suppliers = (
        client.table("company_suppliers")
        .select("supplier_id,supplier_name")
        .eq("company_id", company_id)
        .eq("active", True)
        .execute()
    ).data or []
    source_rows = (
        client.table("company_price_source_rows")
        .select(
            "row_id,raw_description,normalized_name,raw_price,raw_currency,raw_unit,"
            "raw_vat_included,purchase_unit,calculation_unit,conversion_factor,"
            "normalized_unit,result_status,evidence,reason_codes"
        )
        .eq("company_id", company_id)
        .execute()
    ).data or []
    sources = (
        client.table("company_price_sources")
        .select(
            "source_id,source_name,source_kind,source_url,processing_summary,"
            "processed_at,created_at"
        )
        .eq("company_id", company_id)
        .execute()
    ).data or []

    material_by_id = {
        str(row["company_material_id"]): row for row in materials
    }
    supplier_by_id = {str(row["supplier_id"]): row for row in suppliers}
    source_row_by_id = {str(row["row_id"]): row for row in source_rows}
    source_by_id = {str(row["source_id"]): row for row in sources}
    catalog: list[dict] = []
    for offer in offers:
        material = material_by_id.get(str(offer.get("company_material_id")))
        if not material:
            continue
        supplier = supplier_by_id.get(str(offer.get("supplier_id")), {})
        source_row = source_row_by_id.get(str(offer.get("source_row_id")), {})
        source = source_by_id.get(str(offer.get("source_id")), {})
        category = canonical_price_source_category(str(material.get("category") or "Other"))
        department = PRICE_CATALOG_DEPARTMENTS.get(category, "Wood")
        source_summary = source.get("processing_summary") or {}
        supplier_name = str(supplier.get("supplier_name") or "Unknown supplier")
        if source_summary.get("source_origin") == "company_internal":
            supplier_name = "Internal estimate"
        catalog.append(
            {
                **offer,
                "department": department,
                "material_type": category,
                "canonical_name": str(material.get("canonical_name") or "Material"),
                "original_name": str(source_row.get("raw_description") or ""),
                "source_row": source_row,
                "supplier_name": supplier_name,
                "source_name": str(source.get("source_name") or ""),
                "source_kind": str(source.get("source_kind") or ""),
                "source_url": str(source.get("source_url") or ""),
                "updated_at": (
                    offer.get("valid_from")
                    or source.get("processed_at")
                    or offer.get("created_at")
                    or source.get("created_at")
                ),
            }
        )
    return sorted(
        catalog,
        key=lambda row: (
            PRICE_CATALOG_DEPARTMENT_ORDER.get(str(row["department"]), 99),
            str(row["material_type"]).casefold(),
            str(row["canonical_name"]).casefold(),
            str(row["supplier_name"]).casefold(),
        ),
    )


def load_price_source_bytes(access, reference: str | None) -> bytes | None:
    prefix = f"storage://{PRICE_SOURCE_BUCKET}/"
    if not reference or not reference.startswith(prefix):
        return None
    object_path = reference.removeprefix(prefix)
    company_prefix = f"{access.company_id}/"
    if not object_path.startswith(company_prefix):
        raise PermissionError("Price source is not available to this company.")
    client = get_supabase_client()
    assert_company_owner(client, str(access.user_id), str(access.company_id))
    return client.storage.from_(PRICE_SOURCE_BUCKET).download(object_path)


def create_price_source_download_url(
    access,
    reference: str | None,
    *,
    file_name: str,
    expires_in: int = 3600,
) -> str | None:
    """Return a short-lived direct-download URL for an owned source file."""
    prefix = f"storage://{PRICE_SOURCE_BUCKET}/"
    if not reference or not reference.startswith(prefix):
        return None
    object_path = reference.removeprefix(prefix)
    company_prefix = f"{access.company_id}/"
    if not object_path.startswith(company_prefix):
        raise PermissionError("Price source is not available to this company.")
    client = get_supabase_client()
    assert_company_owner(client, str(access.user_id), str(access.company_id))
    response = client.storage.from_(PRICE_SOURCE_BUCKET).create_signed_url(
        object_path,
        expires_in,
        {"download": Path(file_name).name},
    )
    if isinstance(response, dict):
        return response.get("signedURL") or response.get("signed_url")
    return getattr(response, "signedURL", None) or getattr(response, "signed_url", None)


def load_price_source_rows(access, source_id: str) -> list[dict]:
    client = get_supabase_client()
    assert_company_owner(client, str(access.user_id), str(access.company_id))
    return (
        client.table("company_price_source_rows")
        .select("*")
        .eq("company_id", str(access.company_id))
        .eq("source_id", source_id)
        .order("source_row_number")
        .execute()
    ).data or []


def list_unresolved_price_source_rows(access) -> list[dict]:
    """Return reviewable rows from non-archived sources with source context."""
    client = get_supabase_client()
    company_id = str(access.company_id)
    assert_company_owner(client, str(access.user_id), company_id)
    rows = (
        client.table("company_price_source_rows")
        .select("*")
        .eq("company_id", company_id)
        .eq("result_status", "unresolved")
        .order("created_at", desc=True)
        .execute()
    ).data or []
    if not rows:
        return []
    sources = (
        client.table("company_price_sources")
        .select(
            "source_id,source_name,source_kind,source_url,storage_path,mime_type,"
            "category,document_date,currency,status,processing_summary,"
            "company_suppliers(supplier_name)"
        )
        .eq("company_id", company_id)
        .neq("status", "archived")
        .execute()
    ).data or []
    sources_by_id = {str(source["source_id"]): source for source in sources}
    return [
        {**row, "source": sources_by_id[str(row["source_id"])]}
        for row in rows
        if str(row.get("source_id")) in sources_by_id
    ]


def _owned_price_source(client, company_id: str, source_id: str) -> dict:
    rows = (
        client.table("company_price_sources")
        .select("*")
        .eq("company_id", company_id)
        .eq("source_id", source_id)
        .limit(1)
        .execute()
    ).data or []
    if not rows or rows[0].get("status") == "archived":
        raise PriceSourceError("Price source not found")
    return rows[0]


def _refresh_price_source_summary(client, company_id: str, source_id: str) -> None:
    source = _owned_price_source(client, company_id, source_id)
    rows = (
        client.table("company_price_source_rows")
        .select("result_status,raw_vat_included,evidence")
        .eq("company_id", company_id)
        .eq("source_id", source_id)
        .execute()
    ).data or []
    ready = sum(row.get("result_status") in {"new", "updated"} for row in rows)
    new = sum(
        row.get("result_status") == "new"
        and (row.get("evidence") or {}).get("comparison_status") != "unchanged"
        for row in rows
    )
    unchanged = sum(
        row.get("result_status") in {"new", "updated"}
        and (row.get("evidence") or {}).get("comparison_status") == "unchanged"
        for row in rows
    )
    updated = ready - new - unchanged
    unresolved = sum(row.get("result_status") == "unresolved" for row in rows)
    excluded = sum(row.get("result_status") == "excluded" for row in rows)
    material_types = sorted(
        {
            canonical_price_source_category(str((row.get("evidence") or {}).get("material_type") or "Other"))
            for row in rows
            if row.get("result_status") != "excluded"
        },
        key=str.casefold,
    )
    vat_values = {
        row.get("raw_vat_included")
        for row in rows
        if row.get("result_status") != "excluded"
    }
    vat_mode = (
        "included" if vat_values == {True}
        else "excluded" if vat_values == {False}
        else "unknown" if vat_values in (set(), {None})
        else "mixed"
    )
    summary = dict(source.get("processing_summary") or {})
    summary.update(
        {
            "ready": ready,
            "new": new,
            "updated": updated,
            "unchanged": unchanged,
            "unresolved": unresolved,
            "excluded": excluded,
            "total": len(rows),
            "material_types": material_types,
        }
    )
    category = material_types[0] if len(material_types) == 1 else "Mixed"
    client.table("company_price_sources").update(
        {
            "category": category,
            "vat_mode": vat_mode,
            "status": "ready" if unresolved == 0 else "partial",
            "processing_summary": summary,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }
    ).eq("company_id", company_id).eq("source_id", source_id).execute()


def save_price_source_row(access, source_id: str, row_id: str, values: dict) -> dict:
    """Validate a reviewed row, activate its offer, and retain its source evidence."""
    client = get_supabase_client()
    company_id = str(access.company_id)
    assert_company_owner(client, str(access.user_id), company_id)
    source = _owned_price_source(client, company_id, source_id)
    matches = (
        client.table("company_price_source_rows")
        .select("*")
        .eq("company_id", company_id)
        .eq("source_id", source_id)
        .eq("row_id", row_id)
        .limit(1)
        .execute()
    ).data or []
    if not matches:
        raise PriceSourceError("Price row not found")
    row = matches[0]

    normalized_name = str(values.get("normalized_name") or "").strip()
    material_type = canonical_price_source_category(str(values.get("material_type") or ""))
    raw_unit = str(values.get("raw_unit") or "").strip()
    raw_currency = str(values.get("raw_currency") or source.get("currency") or "ILS").strip().upper()
    purchase_unit = str(values.get("purchase_unit") or "").strip()
    calculation_unit = str(values.get("calculation_unit") or "").strip()
    vat_mode = str(values.get("vat_mode") or "unknown")
    try:
        raw_price = float(values.get("raw_price") or 0)
        conversion_factor = float(values.get("conversion_factor") or 0)
    except (TypeError, ValueError) as exc:
        raise PriceSourceError("Enter a valid price and conversion quantity") from exc
    if not normalized_name:
        raise PriceSourceError("Enter a material name")
    if material_type not in PRICE_SOURCE_CATEGORIES:
        raise PriceSourceError("Choose a material type")
    if raw_price <= 0:
        raise PriceSourceError("Enter a positive price")
    if len(raw_currency) != 3:
        raise PriceSourceError("Enter a three-letter currency code")
    if not raw_unit:
        raise PriceSourceError("Enter the source unit")
    if purchase_unit not in CANONICAL_UNIT_CODES or purchase_unit in {"unknown", "other"}:
        raise PriceSourceError("Choose the purchase unit")
    if calculation_unit not in CANONICAL_UNIT_CODES or calculation_unit in {"unknown", "other"}:
        raise PriceSourceError("Choose the calculation unit")
    if conversion_factor <= 0:
        raise PriceSourceError("Enter a positive conversion quantity")
    if vat_mode not in {"included", "excluded"}:
        raise PriceSourceError("Choose whether VAT is included")

    normalized_price = raw_price / conversion_factor
    normalized_key = _normalized_name(normalized_name)
    existing = (
        client.table("company_material_items")
        .select("company_material_id")
        .eq("company_id", company_id)
        .eq("category", material_type)
        .eq("normalized_name", normalized_key)
        .limit(1)
        .execute()
    ).data or []
    if existing:
        material_id = existing[0]["company_material_id"]
        result_status = "updated"
    else:
        material = client.table("company_material_items").insert(
            {
                "company_id": company_id,
                "category": material_type,
                "canonical_name": normalized_name,
                "normalized_name": normalized_key,
                "preferred_unit": calculation_unit,
                "created_from_source_id": source_id,
            }
        ).execute().data[0]
        material_id = material["company_material_id"]
        result_status = "new"

    client.table("company_material_offers").update({"status": "archived"}).eq(
        "company_id", company_id
    ).eq("source_row_id", row_id).neq("status", "archived").execute()
    evidence = dict(row.get("evidence") or {})
    evidence["material_type"] = material_type
    evidence["comparison_status"] = result_status
    resolved_codes = {
        "below_auto_activation_threshold",
        "material_type_unresolved",
        "package_conversion_unresolved",
        "vat_basis_unknown",
    }
    reason_codes = sorted(
        (set(row.get("reason_codes") or []) - resolved_codes) | {"reviewed_by_user"}
    )
    updated = client.table("company_price_source_rows").update(
        {
            "raw_price": raw_price,
            "raw_currency": raw_currency,
            "raw_unit": raw_unit,
            "raw_vat_included": vat_mode == "included",
            "normalized_name": normalized_name,
            "normalized_price": normalized_price,
            "purchase_unit": purchase_unit,
            "calculation_unit": calculation_unit,
            "conversion_factor": conversion_factor,
            "normalized_unit": calculation_unit,
            "conversion_basis": {"description": "Reviewed by company owner"},
            "company_material_id": material_id,
            "result_status": result_status,
            "reason_codes": reason_codes,
            "evidence": evidence,
        }
    ).eq("company_id", company_id).eq("source_id", source_id).eq("row_id", row_id).execute().data[0]
    client.table("company_material_offers").insert(
        {
            "company_id": company_id,
            "company_material_id": material_id,
            "supplier_id": source.get("supplier_id"),
            "source_id": source_id,
            "source_row_id": row_id,
            "supplier_sku": row.get("raw_sku"),
            "source_price": raw_price,
            "source_unit": raw_unit,
            "purchase_unit": purchase_unit,
            "calculation_unit": calculation_unit,
            "conversion_factor": conversion_factor,
            "normalized_price": normalized_price,
            "normalized_unit": calculation_unit,
            "currency": raw_currency,
            "vat_included": vat_mode == "included",
            "valid_from": source.get("document_date"),
            "confidence": row.get("confidence") or 0,
        }
    ).execute()
    if not source.get("currency"):
        client.table("company_price_sources").update({"currency": raw_currency}).eq(
            "company_id", company_id
        ).eq("source_id", source_id).execute()
    _refresh_price_source_summary(client, company_id, source_id)
    return updated


def remove_price_source_row(access, source_id: str, row_id: str) -> None:
    """Remove one offer from active pricing while retaining its audit record."""
    client = get_supabase_client()
    company_id = str(access.company_id)
    assert_company_owner(client, str(access.user_id), company_id)
    _owned_price_source(client, company_id, source_id)
    matches = (
        client.table("company_price_source_rows")
        .select("reason_codes")
        .eq("company_id", company_id)
        .eq("source_id", source_id)
        .eq("row_id", row_id)
        .limit(1)
        .execute()
    ).data or []
    if not matches:
        raise PriceSourceError("Price row not found")
    client.table("company_material_offers").update({"status": "archived"}).eq(
        "company_id", company_id
    ).eq("source_row_id", row_id).neq("status", "archived").execute()
    reason_codes = sorted(set((matches[0].get("reason_codes") or []) + ["removed_by_user"]))
    client.table("company_price_source_rows").update(
        {"result_status": "excluded", "reason_codes": reason_codes}
    ).eq("company_id", company_id).eq("source_id", source_id).eq("row_id", row_id).execute()
    _refresh_price_source_summary(client, company_id, source_id)


def process_price_source(
    access,
    *,
    department: str,
    uploaded_file=None,
    source_url: str = "",
    trace=None,
) -> PriceSourceProcessResult:
    """Process and persist one source. Ambiguous rows stay non-active."""
    process_started = time.perf_counter()
    department = _validate_department(department)
    if (uploaded_file is None) == (not source_url.strip()):
        raise PriceSourceError("Add one file or one supplier URL.")

    company_id = str(access.company_id)
    client = get_supabase_client()
    assert_company_owner(client, str(access.user_id), company_id)
    source_id = str(uuid4())

    if uploaded_file is not None:
        source_read_started = time.perf_counter()
        source_name = str(uploaded_file.name)
        suffix = _source_extension(source_name)
        source_bytes = uploaded_file.getvalue()
        if not source_bytes or len(source_bytes) > MAX_SOURCE_BYTES:
            raise PriceSourceError("The price source must be between 1 byte and 50 MB")
        source_kind = "file"
        resolved_url = None
        _emit_duration(
            trace,
            "server.price_source_read",
            source_read_started,
            source_bytes=len(source_bytes),
        )
        parse_started = time.perf_counter()
        extracted_text = extract_spreadsheet_text(source_name, source_bytes)
        template_sha256 = price_source_template_fingerprint(source_name, source_bytes)
        _emit_duration(
            trace,
            "server.price_source_spreadsheet_parse",
            parse_started,
            extracted_chars=len(extracted_text),
        )
        mime_type = CONTENT_TYPES[suffix]
    else:
        fetch_started = time.perf_counter()
        resolved_url, source_bytes, extracted_text = fetch_public_page(source_url)
        _emit_duration(
            trace,
            "server.price_source_url_fetch",
            fetch_started,
            source_bytes=len(source_bytes),
            extracted_chars=len(extracted_text),
        )
        source_name = resolved_url
        source_kind = "url"
        suffix = ".html"
        mime_type = "text/html"
        template_sha256 = None

    source_digest = sha256(source_bytes).hexdigest()
    duplicate = (
        client.table("company_price_sources")
        .select("source_id,processing_summary")
        .eq("company_id", company_id)
        .eq("source_sha256", source_digest)
        .neq("status", "archived")
        .limit(1)
        .execute()
    ).data or []
    if duplicate:
        return PriceSourceProcessResult(
            source_id=str(duplicate[0]["source_id"]),
            summary=_unchanged_duplicate_summary(duplicate[0]),
        )

    agent_started = time.perf_counter()
    result = run_price_source_agent(
        company_id=company_id,
        department=department,
        source_name=source_name,
        source_bytes=source_bytes if suffix in {".pdf", ".jpg", ".jpeg", ".png"} else None,
        extracted_text=extracted_text,
        import_id=source_id,
        trace=trace,
    )
    guard_price_source_department(result, department)
    material_types = price_source_material_types(result)
    category = material_types[0] if len(material_types) == 1 else "Mixed"
    semantic_sha256 = price_source_semantic_fingerprint(result)
    source_family_sha256, source_family_code = price_source_family_identity(
        result,
        source_name=source_name,
        source_kind=source_kind,
        template_sha256=template_sha256,
    )
    _emit_duration(
        trace,
        "server.price_source_agent",
        agent_started,
        extracted_rows=len(result.get("rows") or []),
    )
    usage_event = result.pop("_agent_usage", None)
    if usage_event:
        usage_started = time.perf_counter()
        try:
            insert_agent_usage_event(client, usage_event)
        except Exception:
            logger.exception("Price source agent usage persistence failed")
        _emit_duration(trace, "server.price_source_usage_persist", usage_started)
    previous_revision, source_revision = _find_previous_source_revision(
        client,
        company_id,
        family_sha256=source_family_sha256,
        semantic_sha256=semantic_sha256,
    )

    benchmark_started = time.perf_counter()
    legacy_materials = (
        client.table("materials")
        .select("material_name,price,price_unit")
        .eq("company_id", company_id)
        .execute()
    ).data or []
    apply_legacy_price_benchmark(result, legacy_materials)
    _emit_duration(
        trace,
        "server.price_source_legacy_benchmark",
        benchmark_started,
        benchmark_rows=len(legacy_materials),
    )
    object_path = f"{company_id}/{source_id}{suffix}"
    storage_started = time.perf_counter()
    client.storage.from_(PRICE_SOURCE_BUCKET).upload(
        object_path,
        source_bytes,
        file_options={"content-type": mime_type, "upsert": "false"},
    )
    _emit_duration(
        trace,
        "server.price_source_storage_upload",
        storage_started,
        source_bytes=len(source_bytes),
    )

    superseded_offer_ids: list[str] = []
    try:
        database_started = time.perf_counter()
        supplier_name = str(result["supplier_name"]).strip()
        supplier_id = None
        if supplier_name:
            existing_suppliers = (
                client.table("company_suppliers")
                .select("supplier_id,categories")
                .eq("company_id", company_id)
                .eq("normalized_name", _normalized_name(supplier_name))
                .limit(1)
                .execute()
            ).data or []
            categories = sorted(
                set((existing_suppliers[0].get("categories") or []) + material_types)
            ) if existing_suppliers else material_types
            supplier_row = (
                client.table("company_suppliers")
                .upsert(
                    {
                        "company_id": company_id,
                        "supplier_name": supplier_name,
                        "normalized_name": _normalized_name(supplier_name),
                        "categories": categories,
                        "updated_at": datetime.now(timezone.utc).isoformat(),
                    },
                    on_conflict="company_id,normalized_name",
                )
                .execute()
            ).data[0]
            supplier_id = supplier_row["supplier_id"]
        ready_count = sum(row["status"] == "ready" for row in result["rows"])
        unresolved_count = sum(row["status"] == "unresolved" for row in result["rows"])
        excluded_count = sum(row["status"] == "excluded" for row in result["rows"])
        new_count = 0
        updated_count = 0
        unchanged_count = 0
        status = "ready" if not unresolved_count else "partial"
        source_summary = {
            "ready": ready_count,
            "new": 0,
            "updated": 0,
            "unchanged": 0,
            "unresolved": unresolved_count,
            "excluded": excluded_count,
            "total": len(result["rows"]),
            "document_number": result["document_number"],
            "source_origin": result["source_origin"],
            "price_context": result["price_context"],
            "document_subtotal": result["document_subtotal"],
            "document_vat_amount": result["document_vat_amount"],
            "document_total": result["document_total"],
            "material_types": material_types,
            "semantic_sha256": semantic_sha256,
            "template_sha256": template_sha256,
            "source_family_sha256": source_family_sha256,
            "source_family_code": source_family_code,
            "source_revision": source_revision,
            "previous_source_id": (
                str(previous_revision.get("source_id")) if previous_revision else None
            ),
            "agent_duration_seconds": (
                usage_event.get("duration_seconds") if usage_event else None
            ),
            "token_cost": usage_event.get("total_cost_usd") if usage_event else None,
            "input_tokens": usage_event.get("input_tokens") if usage_event else None,
            "output_tokens": usage_event.get("output_tokens") if usage_event else None,
            "model": usage_event.get("model") if usage_event else None,
            "prompt_version": usage_event.get("prompt_version") if usage_event else None,
        }
        client.table("company_price_sources").insert(
            {
                "source_id": source_id,
                "company_id": company_id,
                "supplier_id": supplier_id,
                "category": category,
                "source_kind": source_kind,
                "source_name": source_name,
                "source_url": resolved_url,
                "storage_path": f"storage://{PRICE_SOURCE_BUCKET}/{object_path}",
                "mime_type": mime_type,
                "source_sha256": source_digest,
                "document_type": result["document_type"],
                "document_date": _date_or_none(result["document_date"]),
                "currency": result["currency"] or None,
                "vat_mode": result["vat_mode"],
                "status": status,
                "processing_summary": source_summary,
                "created_by": str(access.user_id),
                "processed_at": datetime.now(timezone.utc).isoformat(),
            }
        ).execute()

        active_offers = (
            client.table("company_material_offers")
            .select(
                "offer_id,company_material_id,supplier_id,supplier_sku,source_price,source_unit,"
                "purchase_unit,calculation_unit,conversion_factor,normalized_price,"
                "currency,vat_included,status"
            )
            .eq("company_id", company_id)
            .eq("status", "active")
            .execute()
        ).data or []
        offers_by_identity: dict[tuple[str, str], list[dict]] = {}
        offers_by_sku: dict[tuple[str, str], list[dict]] = {}
        for offer in active_offers:
            identity = (
                str(offer.get("company_material_id") or ""),
                str(offer.get("supplier_id") or ""),
            )
            offers_by_identity.setdefault(identity, []).append(offer)
            sku = _normalized_name(str(offer.get("supplier_sku") or ""))
            if sku:
                offers_by_sku.setdefault((str(offer.get("supplier_id") or ""), sku), []).append(
                    offer
                )

        for row in result["rows"]:
            row_category = canonical_price_source_category(str(row["material_type"]))
            material_id = None
            result_status = row["status"]
            if result_status == "ready":
                normalized = _normalized_name(row["normalized_name"])
                sku = _normalized_name(str(row.get("raw_sku") or ""))
                sku_offers = offers_by_sku.get((str(supplier_id or ""), sku), []) if sku else []
                existing = (
                    [{"company_material_id": sku_offers[0]["company_material_id"]}]
                    if sku_offers
                    else (
                        client.table("company_material_items")
                        .select("company_material_id")
                        .eq("company_id", company_id)
                        .eq("category", row_category)
                        .eq("normalized_name", normalized)
                        .limit(1)
                        .execute()
                    ).data or []
                )
                if existing:
                    material_id = existing[0]["company_material_id"]
                    identity = (str(material_id), str(supplier_id or ""))
                    matching_offers = offers_by_identity.get(identity, [])
                    if matching_offers and any(
                        price_offer_matches_row(
                            offer,
                            row,
                            default_currency=str(result.get("currency") or ""),
                        )
                        for offer in matching_offers
                    ):
                        result_status = "unchanged"
                        unchanged_count += 1
                    else:
                        result_status = "updated"
                        updated_count += 1
                else:
                    material = client.table("company_material_items").insert(
                        {
                            "company_id": company_id,
                            "category": row_category,
                            "canonical_name": row["normalized_name"],
                            "normalized_name": normalized,
                            "preferred_unit": row["calculation_unit"],
                            "created_from_source_id": source_id,
                        }
                    ).execute().data[0]
                    material_id = material["company_material_id"]
                    result_status = "new"
                    new_count += 1

            inserted_row = client.table("company_price_source_rows").insert(
                {
                    "source_id": source_id,
                    "company_id": company_id,
                    "source_row_number": row["source_row_number"],
                    "raw_description": row["raw_description"],
                    "raw_sku": row["raw_sku"] or None,
                    "raw_price": row["raw_price"],
                    "raw_currency": row["raw_currency"] or None,
                    "raw_unit": row["raw_unit"] or None,
                    "raw_package_quantity": row["raw_package_quantity"] or None,
                    "raw_quantity": row["raw_quantity"] or None,
                    "raw_line_total": row["raw_line_total"] or None,
                    "raw_vat_included": True if row["raw_vat_mode"] == "included" else False if row["raw_vat_mode"] == "excluded" else None,
                    "normalized_name": row["normalized_name"] or None,
                    "normalized_price": row["normalized_price"] or None,
                    "purchase_unit": row["purchase_unit"] or None,
                    "calculation_unit": row["calculation_unit"] or None,
                    "conversion_factor": row["conversion_factor"] or None,
                    "normalized_unit": row["calculation_unit"] or None,
                    "conversion_basis": {"description": row["conversion_basis"]},
                    "company_material_id": material_id,
                    "result_status": "updated" if result_status == "unchanged" else result_status,
                    "confidence": row["confidence"],
                    "reason_codes": row["reason_codes"],
                    "evidence": {
                        "reference": row["evidence_reference"],
                        "material_type": row_category,
                        "discount_percent": row["raw_discount_percent"],
                        "discount_amount": row["raw_discount_amount"],
                        "comparison_status": result_status,
                    },
                }
            ).execute().data[0]

            if material_id and result_status in {"new", "updated"}:
                identity = (str(material_id), str(supplier_id or ""))
                for previous_offer in offers_by_identity.get(identity, []):
                    client.table("company_material_offers").update(
                        {"status": "superseded"}
                    ).eq("offer_id", previous_offer["offer_id"]).execute()
                    superseded_offer_ids.append(str(previous_offer["offer_id"]))
                inserted_offer = client.table("company_material_offers").insert(
                    {
                        "company_id": company_id,
                        "company_material_id": material_id,
                        "supplier_id": supplier_id,
                        "source_id": source_id,
                        "source_row_id": inserted_row["row_id"],
                        "supplier_sku": row["raw_sku"] or None,
                        "source_price": row["raw_price"],
                        "source_unit": row["raw_unit"],
                        "purchase_unit": row["purchase_unit"],
                        "calculation_unit": row["calculation_unit"],
                        "conversion_factor": row["conversion_factor"],
                        "normalized_price": row["normalized_price"],
                        "normalized_unit": row["calculation_unit"],
                        "currency": row["raw_currency"] or result["currency"],
                        "vat_included": True if row["raw_vat_mode"] == "included" else False if row["raw_vat_mode"] == "excluded" else None,
                        "valid_from": _date_or_none(result["document_date"]),
                        "confidence": row["confidence"],
                    }
                ).execute().data[0]
                offers_by_identity[identity] = [inserted_offer]
                if sku:
                    offers_by_sku[(str(supplier_id or ""), sku)] = [inserted_offer]

        source_summary.update(
            {
                "new": new_count,
                "updated": updated_count,
                "unchanged": unchanged_count,
            }
        )
        client.table("company_price_sources").update(
            {"processing_summary": source_summary}
        ).eq("company_id", company_id).eq("source_id", source_id).execute()
        _emit_duration(
            trace,
            "server.price_source_database_apply",
            database_started,
            extracted_rows=len(result["rows"]),
            active_rows=ready_count,
            unresolved_rows=unresolved_count,
            excluded_rows=excluded_count,
        )
        _emit_duration(
            trace,
            "server.price_source_total",
            process_started,
            extracted_rows=len(result["rows"]),
        )
        return PriceSourceProcessResult(source_id=source_id, summary=source_summary)
    except Exception:
        cleanup_started = time.perf_counter()
        try:
            client.table("company_material_offers").delete().eq("source_id", source_id).execute()
            for offer_id in superseded_offer_ids:
                client.table("company_material_offers").update({"status": "active"}).eq(
                    "offer_id", offer_id
                ).execute()
            client.table("company_price_source_rows").delete().eq("source_id", source_id).execute()
            client.table("company_material_items").delete().eq("created_from_source_id", source_id).execute()
            client.table("company_price_sources").delete().eq("source_id", source_id).execute()
        except Exception:
            pass
        client.storage.from_(PRICE_SOURCE_BUCKET).remove([object_path])
        _emit_duration(trace, "server.price_source_cleanup", cleanup_started)
        raise
